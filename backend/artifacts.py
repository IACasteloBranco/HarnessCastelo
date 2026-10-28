"""Turn a model response into a downloadable document or spreadsheet."""

import csv
import io
import json
import re
from copy import copy
from datetime import datetime
from itertools import islice
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.utils import ImageReader
from reportlab.platypus import (HRFlowable, ListFlowable, ListItem, Paragraph,
                                Preformatted, SimpleDocTemplate, Spacer, Table,
                                TableStyle)


MIME = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


class InvalidArtifact(ValueError):
    pass


NAVY = colors.HexColor("#102f4a")
GOLD = colors.HexColor("#c9a86b")
INK = colors.HexColor("#183248")
MUTED = colors.HexColor("#526579")
PALE = colors.HexColor("#f3ede1")
LOGO = Path(__file__).resolve().parents[1] / "frontend" / "public" / "castelo-branco.png"


def _inline(text: str) -> str:
    text = escape(text)
    text = re.sub(r"`([^`]+)`", r'<font name="Courier">\1</font>', text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    return re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<i>\1</i>", text)


def _pdf_content(answer: str, width: float) -> tuple[str, list]:
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("ReportTitle", parent=styles["Title"], fontName="Helvetica-Bold",
                              fontSize=23, leading=28, textColor=NAVY, alignment=TA_LEFT,
                              spaceAfter=12))
    styles.add(ParagraphStyle("ReportH1", parent=styles["Heading1"], fontName="Helvetica-Bold",
                              fontSize=19, leading=23, textColor=NAVY, spaceBefore=14, spaceAfter=8,
                              keepWithNext=True))
    styles.add(ParagraphStyle("ReportH2", parent=styles["Heading2"], fontName="Helvetica-Bold",
                              fontSize=15, leading=19, textColor=NAVY, spaceBefore=13, spaceAfter=7,
                              keepWithNext=True))
    styles.add(ParagraphStyle("ReportH3", parent=styles["Heading3"], fontName="Helvetica-Bold",
                              fontSize=11, leading=15, textColor=INK, spaceBefore=9, spaceAfter=4,
                              keepWithNext=True))
    styles.add(ParagraphStyle("ReportBody", parent=styles["BodyText"], fontName="Helvetica",
                              fontSize=10, leading=15, textColor=INK, spaceAfter=8))
    styles.add(ParagraphStyle("ReportCell", parent=styles["BodyText"], fontName="Helvetica",
                              fontSize=8.5, leading=11, textColor=INK))
    styles.add(ParagraphStyle("ReportHeaderCell", parent=styles["ReportCell"],
                              fontName="Helvetica-Bold", textColor=colors.white))
    styles.add(ParagraphStyle("ReportCode", parent=styles["Code"], fontName="Courier",
                              fontSize=8, leading=11, textColor=INK))
    styles.add(ParagraphStyle("ReportQuote", parent=styles["ReportBody"], textColor=MUTED,
                              leftIndent=8, rightIndent=6, spaceAfter=0))

    lines = answer.strip().splitlines()
    title = "Documento do Maximus AI"
    for index, line in enumerate(lines):
        if line.strip():
            heading = re.match(r"^#\s+(.+)$", line.strip())
            if heading:
                title = heading.group(1).strip()[:160]
                lines.pop(index)
            break

    story = [Paragraph(_inline(title), styles["ReportTitle"]),
             HRFlowable(width="100%", thickness=1.5, color=GOLD, spaceAfter=16)]
    paragraph: list[str] = []
    i = 0

    def flush_paragraph() -> None:
        if paragraph:
            story.append(Paragraph(_inline(" ".join(part.strip() for part in paragraph)), styles["ReportBody"]))
            paragraph.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            i += 1
            continue
        if stripped.startswith("```"):
            flush_paragraph()
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            block = Table([[Preformatted(escape("\n".join(code)) or " ", styles["ReportCode"],
                                          maxLineLength=88)]],
                          colWidths=[width])
            block.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#eef2f5")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#d9e0e3")),
                ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]))
            story.extend((block, Spacer(1, 9)))
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(
                r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", lines[i + 1]):
            flush_paragraph()
            rows = [[cell.strip() for cell in stripped.strip("|").split("|")]]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
                i += 1
            columns = max(map(len, rows))
            rows = [row + [""] * (columns - len(row)) for row in rows]
            cells = [[Paragraph(_inline(cell), styles["ReportHeaderCell"] if row_index == 0 else styles["ReportCell"])
                      for cell in row] for row_index, row in enumerate(rows)]
            table = Table(cells, colWidths=[width / columns] * columns, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7f8")]),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d9e0e3")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]))
            story.extend((table, Spacer(1, 10)))
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            style = styles["ReportH1"] if len(heading.group(1)) == 1 else styles[
                "ReportH2" if len(heading.group(1)) == 2 else "ReportH3"]
            story.append(Paragraph(_inline(heading.group(2)), style))
            i += 1
            continue
        if re.match(r"^(?:[-*_]\s*){3,}$", stripped):
            flush_paragraph()
            story.extend((HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#d9e0e3")), Spacer(1, 8)))
            i += 1
            continue
        if stripped.startswith(">"):
            flush_paragraph()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            block = Table([[Paragraph(_inline(" ".join(quote)), styles["ReportQuote"])]], colWidths=[width])
            block.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), PALE), ("LINEBEFORE", (0, 0), (0, 0), 2, GOLD),
                ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]))
            story.extend((block, Spacer(1, 9)))
            continue
        list_match = re.match(r"^\s*(?:([-*+])|(\d+)[.)])\s+(.+)$", line)
        if list_match:
            flush_paragraph()
            ordered = list_match.group(2) is not None
            items = []
            while i < len(lines):
                item = re.match(r"^\s*(?:([-*+])|(\d+)[.)])\s+(.+)$", lines[i])
                if not item or (item.group(2) is not None) != ordered:
                    break
                items.append(ListItem(Paragraph(_inline(item.group(3)), styles["ReportBody"]),
                                      leftIndent=12))
                i += 1
            list_options = {"bulletType": "1", "start": int(list_match.group(2) or 1)} if ordered else {
                "bulletType": "bullet"}
            story.extend((ListFlowable(items, leftIndent=20, bulletFontName="Helvetica",
                                       bulletFontSize=8, **list_options), Spacer(1, 4)))
            continue
        paragraph.append(stripped)
        i += 1
    flush_paragraph()
    return title, story


def _pdf_page(canvas, document) -> None:
    page_width, page_height = A4
    canvas.saveState()
    canvas.setFillColor(NAVY)
    canvas.rect(0, page_height - 68, page_width, 68, fill=1, stroke=0)
    canvas.setFillColor(GOLD)
    canvas.rect(0, page_height - 70, page_width, 2, fill=1, stroke=0)
    if LOGO.is_file():
        image = ImageReader(str(LOGO))
        image_width, image_height = image.getSize()
        height = 32
        canvas.drawImage(image, 42, page_height - 52, width=height * image_width / image_height,
                         height=height, preserveAspectRatio=True, mask="auto")
    canvas.setFillColor(colors.HexColor("#e8dcc4"))
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(page_width - 42, page_height - 31,
                           datetime.now().astimezone().strftime("%d/%m/%Y - %H:%M"))
    canvas.setStrokeColor(colors.HexColor("#d9e0e3"))
    canvas.setLineWidth(0.5)
    canvas.line(42, 38, page_width - 42, 38)
    canvas.setFillColor(MUTED)
    canvas.setFont("Helvetica", 8)
    canvas.drawString(42, 24, "Documento gerado pelo Maximus AI")
    canvas.drawRightString(page_width - 42, 24, str(document.page))
    canvas.restoreState()


def _table(source: str) -> tuple[list[str], list[list[str]]]:
    source = source.strip()
    if source.startswith("```"):
        source = source.partition("\n")[2].rsplit("```", 1)[0].strip()
    try:
        payload = json.loads(source)
    except json.JSONDecodeError as exc:
        raise InvalidArtifact("A IA não retornou uma tabela válida. Tente novamente.") from exc
    if not isinstance(payload, dict) or set(payload) != {"columns", "rows"}:
        raise InvalidArtifact("A IA não retornou colunas e linhas válidas.")
    columns, rows = payload["columns"], payload["rows"]
    if (not isinstance(columns, list) or not 1 <= len(columns) <= 50 or
            not all(isinstance(value, str) and value.strip() and len(value) <= 255 for value in columns) or
            not isinstance(rows, list) or len(rows) > 10_000):
        raise InvalidArtifact("A tabela excedeu os limites ou veio incompleta.")
    clean = []
    for row in rows:
        if not isinstance(row, list) or len(row) != len(columns) or any(
            value is not None and not isinstance(value, (str, int, float, bool)) for value in row
        ):
            raise InvalidArtifact("A IA retornou uma linha inválida.")
        clean.append(["" if value is None else str(value) for value in row])
    return columns, clean


def build_artifact(format: str, answer: str) -> bytes:
    if format == "pdf":
        if not answer.strip():
            raise InvalidArtifact("A resposta está vazia.")
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, title="Maximus AI", author="Maximus AI", pagesize=A4,
                                leftMargin=46, rightMargin=46, topMargin=92, bottomMargin=52)
        title, story = _pdf_content(answer, doc.width)
        doc.title = title
        doc.subject = "Documento gerado pelo Maximus AI"
        doc.build(story, onFirstPage=_pdf_page, onLaterPages=_pdf_page)
        data = output.getvalue()
    elif format in ("csv", "xlsx"):
        columns, rows = _table(answer)
        if format == "csv":
            output = io.StringIO(newline="")
            writer = csv.writer(output)
            writer.writerows([["'" + value if value.startswith(("=", "+", "-", "@")) else value
                              for value in row] for row in [columns, *rows]])
            data = ("\ufeff" + output.getvalue()).encode("utf-8")
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Dados"
            for row in [columns, *rows]:
                sheet.append(["'" + value if value.startswith(("=", "+", "-", "@")) else value for value in row])
            for cell in sheet[1]:
                font = copy(cell.font)
                font.bold = True
                font.color = "FFFFFF"
                cell.font = font
                cell.fill = PatternFill("solid", fgColor="102F4A")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = Border(bottom=Side(style="medium", color="C9A86B"))
            for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=1):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if row_index % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor="F3F6F8")
            for column in sheet.columns:
                width = min(42, max(10, max(len(str(cell.value or "")) for cell in islice(column, 301)) + 2))
                sheet.column_dimensions[column[0].column_letter].width = width
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_view.showGridLines = False
            workbook.properties.creator = "Maximus AI"
            output = io.BytesIO()
            workbook.save(output)
            data = output.getvalue()
    else:
        raise InvalidArtifact("Formato não suportado.")
    if len(data) > 20_000_000:
        raise InvalidArtifact("O arquivo gerado excedeu 20 MB.")
    return data


def preview_table(format: str, data: bytes) -> dict:
    workbook = None
    if format in ("csv", "tsv"):
        rows = csv.reader(io.StringIO(data.decode("utf-8-sig")), delimiter="\t" if format == "tsv" else ",")
    elif format == "xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
    else:
        raise InvalidArtifact("Formato não suportado.")
    header = next(rows, ())
    preview = [["" if value is None else str(value) for value in row] for row in islice(rows, 101)]
    if workbook:
        workbook.close()
    return {"columns": ["" if value is None else str(value) for value in header],
            "rows": preview[:100], "truncated": len(preview) > 100}
